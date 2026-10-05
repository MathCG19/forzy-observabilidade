"""Funções de formatação ABNT (NBR 14724, 6024, 10520) sobre python-docx."""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONTE = "Arial"
PRETO = RGBColor(0, 0, 0)


def _fonte(estilo_ou_run, tamanho=12, negrito=None, caixa_alta=None):
    f = estilo_ou_run.font
    f.name = FONTE
    f.size = Pt(tamanho)
    f.color.rgb = PRETO
    if negrito is not None:
        f.bold = negrito
    if caixa_alta is not None:
        f.all_caps = caixa_alta
    rpr = estilo_ou_run.element.get_or_add_rPr() if hasattr(estilo_ou_run.element, "get_or_add_rPr") \
        else estilo_ou_run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(atributo), FONTE)
    for atributo in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(atributo)) is not None:
            del rfonts.attrib[qn(atributo)]


def _paragrafo_fmt(pf, entrelinha=1.5, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=Cm(1.25),
                   antes=Pt(0), depois=Pt(0)):
    pf.line_spacing = entrelinha
    pf.alignment = alinhamento
    pf.first_line_indent = recuo
    pf.left_indent = Cm(0)
    pf.space_before = antes
    pf.space_after = depois


def novo_documento() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.left_margin = Cm(3), Cm(3)
    sec.bottom_margin, sec.right_margin = Cm(2), Cm(2)
    sec.header_distance = Cm(2)
    sec.footer_distance = Cm(1.25)

    estilos = doc.styles
    normal = estilos["Normal"]
    _fonte(normal)
    _paragrafo_fmt(normal.paragraph_format)

    for nome, nivel, negrito, caixa in (("Heading 1", 1, True, True), ("Heading 2", 2, False, True),
                                        ("Heading 3", 3, True, False)):
        h = estilos[nome]
        _fonte(h, 12, negrito, caixa)
        h.font.italic = False
        _paragrafo_fmt(h.paragraph_format, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, recuo=Cm(0),
                       antes=Pt(0) if nivel == 1 else Pt(18), depois=Pt(18))
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.page_break_before = nivel == 1

    def estilo(nome, base="Normal"):
        s = estilos.add_style(nome, 1)
        s.base_style = estilos[base]
        return s

    titulo = estilo("Titulo sem numero")
    _fonte(titulo, 12, True, True)
    _paragrafo_fmt(titulo.paragraph_format, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, recuo=Cm(0), depois=Pt(18))

    legenda = estilo("Legenda ABNT")
    _fonte(legenda, 10)
    _paragrafo_fmt(legenda.paragraph_format, 1.0, WD_ALIGN_PARAGRAPH.CENTER, Cm(0), antes=Pt(12), depois=Pt(4))
    legenda.paragraph_format.keep_with_next = True

    fonte = estilo("Fonte ABNT")
    _fonte(fonte, 10)
    _paragrafo_fmt(fonte.paragraph_format, 1.0, WD_ALIGN_PARAGRAPH.CENTER, Cm(0), antes=Pt(4), depois=Pt(12))

    ref = estilo("Referencia")
    _fonte(ref)
    _paragrafo_fmt(ref.paragraph_format, 1.0, WD_ALIGN_PARAGRAPH.LEFT, Cm(0), depois=Pt(12))

    celula = estilo("Celula")
    _fonte(celula, 10)
    _paragrafo_fmt(celula.paragraph_format, 1.0, WD_ALIGN_PARAGRAPH.LEFT, Cm(0))

    lista = estilo("Item de lista")
    _paragrafo_fmt(lista.paragraph_format, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, recuo=Cm(0))

    # Word pergunta se quer atualizar os campos (sumário e listas) ao abrir o arquivo.
    atualizar = OxmlElement("w:updateFields")
    atualizar.set(qn("w:val"), "true")
    doc.settings.element.append(atualizar)
    return doc


def paragrafo(doc, texto="", estilo="Normal", negrito=False, alinhamento=None, recuo=None, tamanho=None):
    p = doc.add_paragraph(style=estilo)
    if texto:
        run = p.add_run(texto)
        run.bold = negrito or None
        if tamanho:
            run.font.size = Pt(tamanho)
    if alinhamento is not None:
        p.paragraph_format.alignment = alinhamento
    if recuo is not None:
        p.paragraph_format.first_line_indent = recuo
    return p


def texto_com_italico(doc, partes, estilo="Normal"):
    """partes: lista de (texto, italico). Usado para termos estrangeiros."""
    p = doc.add_paragraph(style=estilo)
    for texto, italico in partes:
        run = p.add_run(texto)
        run.italic = italico or None
    return p


def quebra_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _campo(paragrafo_xml_run_parent, instrucao: str, resultado: str):
    """Insere um campo simples (PAGE, SEQ) com resultado em cache no parágrafo."""
    p = paragrafo_xml_run_parent

    def run_com(elemento):
        r = OxmlElement("w:r")
        r.append(elemento)
        p._p.append(r)
        return r

    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    run_com(inicio)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instrucao} "
    run_com(instr)
    separa = OxmlElement("w:fldChar")
    separa.set(qn("w:fldCharType"), "separate")
    run_com(separa)
    p.add_run(resultado)
    fim = OxmlElement("w:fldChar")
    fim.set(qn("w:fldCharType"), "end")
    run_com(fim)


def campo_indice(doc, instrucao: str, entradas: list[tuple[int, str]]):
    """Campo TOC com as entradas já escritas (sem número de página). F9 no Word completa."""
    if not entradas:
        entradas = [(1, "Atualize o campo (F9) no Word.")]
    paragrafos = []
    for nivel, texto in entradas:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.5 * (nivel - 1))
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragrafos.append((p, texto, nivel))

    primeiro = paragrafos[0][0]
    for tipo, texto in (("begin", None), ("instr", instrucao), ("separate", None)):
        r = OxmlElement("w:r")
        if tipo == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {texto} "
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        r.append(el)
        primeiro._p.append(r)
    for p, texto, nivel in paragrafos:
        run = p.add_run(texto)
        run.bold = (nivel == 1 and instrucao.startswith("TOC \\o")) or None
    r = OxmlElement("w:r")
    fim = OxmlElement("w:fldChar")
    fim.set(qn("w:fldCharType"), "end")
    r.append(fim)
    paragrafos[-1][0]._p.append(r)


class Contador:
    def __init__(self):
        self.n = {"Figura": 0, "Quadro": 0, "Tabela": 0}
        self.entradas = {"Figura": [], "Quadro": [], "Tabela": []}

    def legenda(self, doc, tipo: str, titulo: str):
        self.n[tipo] += 1
        numero = self.n[tipo]
        p = doc.add_paragraph(style="Legenda ABNT")
        p.add_run(f"{tipo} ")
        _campo(p, f"SEQ {tipo} \\* ARABIC", str(numero))
        p.add_run(f" - {titulo}")
        self.entradas[tipo].append((1, f"{tipo} {numero} - {titulo}"))
        return numero


def fonte(doc, texto="Fonte: elaborado pelos autores (2026)."):
    return paragrafo(doc, texto, "Fonte ABNT")


def figura(doc, contador: Contador, caminho, titulo: str, largura_cm=15.5, texto_fonte=None):
    numero = contador.legenda(doc, "Figura", titulo)
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(caminho), width=Cm(largura_cm))
    fonte(doc, texto_fonte or "Fonte: elaborado pelos autores (2026).")
    return numero


def _bordas(tabela, fechada: bool):
    """Quadro: todas as bordas. Tabela (IBGE): só linhas horizontais externas e do cabeçalho."""
    tbl_pr = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    lados = ("top", "left", "bottom", "right", "insideH", "insideV") if fechada else ("top", "bottom")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        if lado in lados:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "6")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        bordas.append(el)
    posterior = next((tbl_pr.find(qn(f"w:{t}")) for t in ("shd", "tblLayout", "tblCellMar", "tblLook")
                      if tbl_pr.find(qn(f"w:{t}")) is not None), None)
    if posterior is not None:
        posterior.addprevious(bordas)
    else:
        tbl_pr.append(bordas)
    if not fechada:
        for celula in tabela.rows[0].cells:
            tc_pr = celula._tc.get_or_add_tcPr()
            b = OxmlElement("w:tcBorders")
            el = OxmlElement("w:bottom")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "6")
            el.set(qn("w:color"), "000000")
            b.append(el)
            tc_pr.append(b)


def _repetir_cabecalho(linha):
    tr_pr = linha._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def grade(doc, contador: Contador, tipo: str, titulo: str, cabecalho: list[str], linhas: list[list],
          larguras_cm: list[float] | None = None, texto_fonte=None):
    contador.legenda(doc, tipo, titulo)
    tabela = doc.add_table(rows=1, cols=len(cabecalho))
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = False
    for i, texto in enumerate(cabecalho):
        cel = tabela.rows[0].cells[i]
        cel.paragraphs[0].style = doc.styles["Celula"]
        cel.paragraphs[0].add_run(texto).bold = True
    _repetir_cabecalho(tabela.rows[0])
    for linha in linhas:
        cells = tabela.add_row().cells
        for i, valor in enumerate(linha):
            cells[i].paragraphs[0].style = doc.styles["Celula"]
            cells[i].paragraphs[0].add_run("" if valor is None else str(valor))
    if larguras_cm:
        for linha in tabela.rows:
            for i, largura in enumerate(larguras_cm):
                linha.cells[i].width = Cm(largura)
    _bordas(tabela, fechada=(tipo == "Quadro"))
    fonte(doc, texto_fonte or "Fonte: elaborado pelos autores (2026).")
    return tabela


def destaque_amarelo(doc, linhas: list[str], negrito_primeira=True):
    for i, texto in enumerate(linhas):
        p = doc.add_paragraph(style="Normal")
        if i:
            p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(texto)
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        run.bold = (i == 0 and negrito_primeira) or None


def nova_secao(doc, inicio_numeracao: int | None = None, mostrar_numero: bool = False):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.header.is_linked_to_previous = False
    cab = sec.header.paragraphs[0]
    for run in list(cab.runs):
        run._r.getparent().remove(run._r)
    if mostrar_numero:
        cab.style = doc.styles["Normal"]
        cab.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cab.paragraph_format.first_line_indent = Cm(0)
        _campo(cab, "PAGE", "")
        for run in cab.runs:
            run.font.size = Pt(10)
    # add_section clona as propriedades da seção anterior; tira o reinício herdado.
    for herdado in sec._sectPr.findall(qn("w:pgNumType")):
        sec._sectPr.remove(herdado)
    if inicio_numeracao is not None:
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), str(inicio_numeracao))
        posterior = next((sec._sectPr.find(qn(f"w:{t}")) for t in ("cols", "formProt", "vAlign", "noEndnote",
                          "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid")
                          if sec._sectPr.find(qn(f"w:{t}")) is not None), None)
        if posterior is not None:
            posterior.addprevious(pg)
        else:
            sec._sectPr.append(pg)
    return sec


__all__ = [
    "Cm", "Pt", "WD_ALIGN_PARAGRAPH", "WD_LINE_SPACING", "Contador", "campo_indice", "destaque_amarelo",
    "figura", "fonte", "grade", "nova_secao", "novo_documento", "paragrafo", "quebra_pagina", "texto_com_italico",
]
